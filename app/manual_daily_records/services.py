"""Self-entry and manager bulk import for manual daily records."""
import base64
import binascii
import hmac
from copy import copy
from decimal import Decimal, InvalidOperation
from io import BytesIO

from flask_smorest import abort
from openpyxl import Workbook, load_workbook
from sqlalchemy import tuple_

from datetime import datetime, timezone

from app.extensions import db
from app.cohorts.models import CohortMembership
from app.manual_daily_records.models import (
    ManualDailyRecord,
    ManualImportBatch,
    ManualImportChunk,
)
from app.users.hierarchy import manager_team_user_ids
from app.users.models import User

MANUAL_UPLOAD_HEADERS = (
    "Email",
    "Coder",
    "Production count today",
    "Tech Issues/Downtime (in hrs)",
    "No Inventory/Idle Time (in hrs)",
    "Leave (L) (in hrs)",
    "Meeting/Huddles/Employee Engagement activities (in hrs)",
    "Holiday (H) (in hrs)",
)

MANUAL_MTD_UPLOAD_HEADERS = (
    "Date",
    "Coder",
    "Production count today",
    "Tech Issues/Downtime (in hrs)",
    "No Inventory/Idle Time (in hrs)",
    "Leave (L) (in hrs)",
    "Meeting/Huddles/Employee Engagement activities (in hrs)",
    "Holiday (H) (in hrs)",
)

_ENTRY_FIELDS = (
    "tech_issues_downtime_hours",
    "no_inventory_idle_time_hours",
    "leave_hours",
    "meeting_engagement_hours",
)


def build_manual_upload_template_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Manual Production"
    sheet.append(MANUAL_MTD_UPLOAD_HEADERS)
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    widths = (38, 28, 24, 32, 35, 22, 58, 24)
    for column, width in zip("ABCDEFGH", widths):
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _normalized_name(value):
    return " ".join(str(value or "").strip().casefold().split())


def _decimal_value(value, *, row_number, label, integer=False):
    if value is None or str(value).strip() == "":
        return Decimal("0"), None
    try:
        parsed = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None, {"row": row_number, "message": f"{label} must be a number."}
    if not parsed.is_finite() or parsed < 0:
        return None, {"row": row_number, "message": f"{label} cannot be negative."}
    if integer and parsed != parsed.to_integral_value():
        return None, {"row": row_number, "message": f"{label} must be a whole number."}
    if not integer and parsed > Decimal("10"):
        return None, {"row": row_number, "message": f"{label} cannot exceed 10 hours."}
    return parsed, None


def _decode_workbook(file_base64, source_filename):
    if not source_filename.lower().endswith(".xlsx"):
        abort(400, message="Upload an .xlsx workbook.")
    try:
        payload = base64.b64decode(file_base64, validate=True)
    except (binascii.Error, ValueError):
        abort(400, message="The uploaded workbook payload is not valid base64.")
    if not payload:
        abort(400, message="The uploaded workbook is empty.")
    try:
        return load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception:
        abort(400, message="The uploaded file is not a readable .xlsx workbook.")


def _read_bulk_rows(workbook, uploaded_by_id):
    sheet = workbook.worksheets[0]
    header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = {" ".join(str(value or "").strip().lower().split()): index for index, value in enumerate(header_values)}
    missing = [header for header in MANUAL_UPLOAD_HEADERS if " ".join(header.lower().split()) not in headers]
    if missing:
        abort(400, message=f"The first worksheet is missing required columns: {', '.join(missing)}.")

    team_ids = manager_team_user_ids(uploaded_by_id)
    team_users = User.query.filter(User.id.in_(team_ids), User.is_active.is_(True)).all() if team_ids else []
    users_by_email = {user.email.strip().casefold(): user for user in team_users}
    parsed_rows = []
    row_errors = []
    seen_emails = set()

    def cell(row, header):
        index = headers[" ".join(header.lower().split())]
        return row[index] if index < len(row) else None

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        values = [cell(row, header) for header in MANUAL_UPLOAD_HEADERS]
        if all(value is None or str(value).strip() == "" for value in values):
            continue

        email = str(values[0] or "").strip()
        coder_name = str(values[1] or "").strip()
        row_prefix = {"row": row_number, **({"email": email} if email else {})}
        if not email:
            row_errors.append({**row_prefix, "message": "Email is required."})
            continue
        email_key = email.casefold()
        if email_key in seen_emails:
            row_errors.append({**row_prefix, "message": "Email appears more than once in the workbook."})
            continue
        seen_emails.add(email_key)

        user = users_by_email.get(email_key)
        if user is None:
            row_errors.append({**row_prefix, "message": "Email does not match an active user in your team."})
            continue
        if not coder_name:
            row_errors.append({**row_prefix, "message": "Coder name is required."})
            continue
        expected_name = f"{user.first_name} {user.last_name}"
        if _normalized_name(coder_name) != _normalized_name(expected_name):
            row_errors.append(
                {**row_prefix, "message": f"Coder name does not match {expected_name} for this email."}
            )
            continue

        numeric_specs = (
            (values[2], "Production count today", True),
            (values[3], "Tech Issues/Downtime", False),
            (values[4], "No Inventory/Idle Time", False),
            (values[5], "Leave", False),
            (values[6], "Meeting/Huddles/Employee Engagement activities", False),
            (values[7], "Holiday", False),
        )
        parsed_numbers = []
        numeric_error = None
        for raw_value, label, integer in numeric_specs:
            parsed, numeric_error = _decimal_value(
                raw_value, row_number=row_number, label=label, integer=integer
            )
            if numeric_error:
                numeric_error.update(row_prefix)
                row_errors.append(numeric_error)
                break
            parsed_numbers.append(parsed)
        if numeric_error:
            continue
        if parsed_numbers[5] != 0:
            row_errors.append({**row_prefix, "message": "Holiday must be 0 for manual production uploads."})
            continue

        parsed_rows.append(
            {
                "user": user,
                "production_count": int(parsed_numbers[0]),
                "tech_issues_downtime_hours": parsed_numbers[1],
                "no_inventory_idle_time_hours": parsed_numbers[2],
                "leave_hours": parsed_numbers[3],
                "meeting_engagement_hours": parsed_numbers[4],
            }
        )

    if row_errors:
        abort(
            422,
            message="Nothing was imported. Fix the workbook rows and upload again.",
            data={"rowErrors": row_errors},
        )
    if not parsed_rows:
        abort(400, message="The first worksheet does not contain any data rows.")
    return sheet.title, parsed_rows


def import_manual_daily_records(file_base64, source_filename, record_date, uploaded_by_id):
    workbook = _decode_workbook(file_base64, source_filename)
    sheet_name, rows = _read_bulk_rows(workbook, uploaded_by_id)
    created_count = 0
    updated_count = 0
    positive_user_ids = set()

    for data in rows:
        user = data.pop("user")
        record = ManualDailyRecord.query.filter_by(user_id=user.id, record_date=record_date).first()
        if record is None:
            record = ManualDailyRecord(user_id=user.id, record_date=record_date)
            db.session.add(record)
            created_count += 1
        else:
            updated_count += 1

        record.pvp_count = data["production_count"]
        record.foundation_count = 0
        record.production_count = data["production_count"]
        for field in _ENTRY_FIELDS:
            setattr(record, field, data[field])
        record.status = "pending"
        record.reviewed_by_id = None
        record.reviewed_at = None
        record.rejection_reason = None
        if record.production_count > 0:
            positive_user_ids.add(user.id)

    db.session.flush()
    if positive_user_ids:
        from app.cohorts.services import compute_user_stage_periods

        memberships = CohortMembership.query.filter(CohortMembership.user_id.in_(positive_user_ids)).all()
        for membership in memberships:
            compute_user_stage_periods(membership)
    db.session.commit()
    return {
        "record_date": record_date,
        "source_filename": source_filename,
        "sheet_name": sheet_name,
        "row_count": len(rows),
        "imported_count": len(rows),
        "created_count": created_count,
        "updated_count": updated_count,
    }


def start_manual_import(source_filename, file_checksum, total_rows, uploaded_by_id):
    batch = ManualImportBatch(
        source_filename=source_filename,
        file_checksum=file_checksum.lower(),
        total_rows=total_rows,
        uploaded_by_id=uploaded_by_id,
        status="uploading",
    )
    db.session.add(batch)
    db.session.commit()
    return batch


def _manual_values_match(record, row):
    return (
        record.production_count == row["production_count"]
        and record.tech_issues_downtime_hours == row["tech_issues_downtime_hours"]
        and record.no_inventory_idle_time_hours == row["no_inventory_idle_time_hours"]
        and record.leave_hours == row["leave_hours"]
        and record.meeting_engagement_hours == row["meeting_engagement_hours"]
    )


def process_manual_import_chunk(batch_id, chunk_number, checksum, rows, uploaded_by_id):
    """Upsert one manager-scoped chunk without touching identical records."""
    batch = db.session.get(ManualImportBatch, batch_id)
    if batch is None:
        abort(404, message="Manual production import was not found.")
    if batch.uploaded_by_id != uploaded_by_id:
        abort(403, message="This manual production import belongs to another manager.")
    if batch.status == "completed":
        abort(409, message="This manual production import is already complete.")

    existing_chunk = ManualImportChunk.query.filter_by(
        batch_id=batch_id, chunk_number=chunk_number
    ).first()
    if existing_chunk is not None:
        if not hmac.compare_digest(existing_chunk.checksum, checksum.lower()):
            abort(409, message="This chunk number was already uploaded with different content.")
        return batch, existing_chunk

    if batch.processed_count + len(rows) > batch.total_rows:
        abort(409, message="This chunk would exceed the import's declared row count.")

    team_ids = set(manager_team_user_ids(uploaded_by_id))
    requested_user_ids = {row["user_id"] for row in rows}
    if not requested_user_ids.issubset(team_ids):
        abort(403, message="One or more manual production rows are outside your team.")

    users = User.query.filter(User.id.in_(requested_user_ids)).all()
    users_by_id = {user.id: user for user in users}
    if len(users_by_id) != len(requested_user_ids):
        abort(422, message="One or more users in the manual production upload no longer exist.")

    seen_keys = set()
    for row in rows:
        key = (row["user_id"], row["record_date"])
        if key in seen_keys:
            abort(422, message="A user and date may appear only once in each upload chunk.")
        seen_keys.add(key)
        user = users_by_id[row["user_id"]]
        if not user.is_active and user.last_working_day and row["record_date"] > user.last_working_day:
            abort(
                422,
                message=(
                    f"{user.first_name} {user.last_name} has a last working day of "
                    f"{user.last_working_day.isoformat()}; later records cannot be imported."
                ),
            )

    existing_records = {
        (record.user_id, record.record_date): record
        for record in ManualDailyRecord.query.filter(
            tuple_(ManualDailyRecord.user_id, ManualDailyRecord.record_date).in_(list(seen_keys))
        ).all()
    }
    counts = {"created": 0, "updated": 0, "unchanged": 0}
    positive_user_ids = set()
    for row in rows:
        key = (row["user_id"], row["record_date"])
        record = existing_records.get(key)
        if record is not None and _manual_values_match(record, row):
            counts["unchanged"] += 1
            continue
        if record is None:
            record = ManualDailyRecord(user_id=row["user_id"], record_date=row["record_date"])
            db.session.add(record)
            existing_records[key] = record
            counts["created"] += 1
        else:
            counts["updated"] += 1

        record.pvp_count = row["production_count"]
        record.foundation_count = 0
        record.production_count = row["production_count"]
        for field in _ENTRY_FIELDS:
            setattr(record, field, row[field])
        record.status = "pending"
        record.reviewed_by_id = None
        record.reviewed_at = None
        record.rejection_reason = None
        if record.production_count > 0:
            positive_user_ids.add(record.user_id)

    chunk = ManualImportChunk(
        batch_id=batch.id,
        chunk_number=chunk_number,
        checksum=checksum.lower(),
        row_count=len(rows),
        created_count=counts["created"],
        updated_count=counts["updated"],
        unchanged_count=counts["unchanged"],
    )
    db.session.add(chunk)
    batch.processed_count += len(rows)
    batch.created_count += counts["created"]
    batch.updated_count += counts["updated"]
    batch.unchanged_count += counts["unchanged"]
    db.session.flush()

    if positive_user_ids:
        from app.cohorts.services import compute_user_stage_periods

        memberships = CohortMembership.query.filter(CohortMembership.user_id.in_(positive_user_ids)).all()
        for membership in memberships:
            compute_user_stage_periods(membership)
    db.session.commit()
    return batch, chunk


def complete_manual_import(batch_id, uploaded_by_id):
    batch = db.session.get(ManualImportBatch, batch_id)
    if batch is None:
        abort(404, message="Manual production import was not found.")
    if batch.uploaded_by_id != uploaded_by_id:
        abort(403, message="This manual production import belongs to another manager.")
    if batch.processed_count != batch.total_rows:
        abort(
            409,
            message=f"Upload is incomplete: processed {batch.processed_count} of {batch.total_rows} rows.",
        )
    batch.status = "completed"
    batch.completed_at = datetime.now(timezone.utc)
    db.session.commit()
    return batch


def upsert_own_record(user_id, data):
    """Creates today's-or-any-day's record for `user_id`, or updates it if
    one already exists for that (user, date) pair - unique per §6.1, so
    re-submitting the same day always edits in place rather than creating
    a duplicate.

    Any edit - whether the record was previously Pending, Approved, or
    Rejected - resets it to Pending and clears the prior review: an
    approval or rejection was a decision about these exact numbers, and
    that decision no longer applies once the numbers change (§6.3, applied
    symmetrically to rejection per §9's resubmission assumption).
    """
    record = ManualDailyRecord.query.filter_by(user_id=user_id, record_date=data["record_date"]).first()
    if record is None:
        record = ManualDailyRecord(user_id=user_id, record_date=data["record_date"])
        db.session.add(record)

    pvp_count = data.get("pvp_count")
    foundation_count = data.get("foundation_count")
    if pvp_count is None and foundation_count is None:
        pvp_count, foundation_count = data.get("production_count", 0), 0
    record.pvp_count = pvp_count or 0
    record.foundation_count = foundation_count or 0
    record.production_count = record.pvp_count + record.foundation_count

    for field in _ENTRY_FIELDS:
        setattr(record, field, data[field])

    record.status = "pending"
    record.reviewed_by_id = None
    record.reviewed_at = None
    record.rejection_reason = None

    db.session.commit()

    # A user's first real production day is the authoritative M1 start for
    # future cohorts. Recompute from live facts so a lead and an employee
    # progress identically and no scheduled refresh is required.
    if record.production_count > 0:
        from app.cohorts.models import CohortMembership
        from app.cohorts.services import compute_user_stage_periods

        membership = CohortMembership.query.filter_by(user_id=user_id).first()
        if membership is not None:
            compute_user_stage_periods(membership)
            db.session.commit()
    return record


def approve_record(record, reviewed_by_id):
    """Mutates `record` in place to Approved - does not commit, so a caller
    reviewing several records (the Reports dashboard's bulk-approve) can
    apply this to each one and commit exactly once."""
    record.status = "approved"
    record.reviewed_by_id = reviewed_by_id
    record.reviewed_at = datetime.now(timezone.utc)
    record.rejection_reason = None


def reject_record(record, reviewed_by_id, reason=None):
    """Mutates `record` in place to Rejected - see approve_record() re: no commit."""
    record.status = "rejected"
    record.reviewed_by_id = reviewed_by_id
    record.reviewed_at = datetime.now(timezone.utc)
    record.rejection_reason = reason
