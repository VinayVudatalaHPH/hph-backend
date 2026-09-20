import base64
import binascii
import re
from datetime import date, datetime, time, timedelta
from io import BytesIO

from flask_smorest import abort
from openpyxl import load_workbook

from app.extensions import db
from app.cohorts.models import Cohort, CohortMembership
from app.login_hours.models import LoginHourRecord, LoginHoursUploadBatch
from app.roles.models import Role, RoleType
from app.users.models import Project, User
from app.users.hierarchy import manager_team_user_ids


REQUIRED_HEADERS = {
    "date",
    "employee",
    "firstin",
    "lastout",
    "totalinside",
    "totaloutside",
    "totalspan",
}

# Explicit spelling/order differences confirmed by the supplied Book6
# name/email reference. These are intentionally not fuzzy matches.
ATTENDANCE_NAME_ALIASES = {
    "ajaymishra": "ajayamishra",
    "navyadupati": "navyasridupati",
    "prameelavedula": "prameelavedulla",
    "prathyushapuli": "pratyushapuli",
    "ratnakumarinimmagadda": "rathnakumarinimmagadda",
    "rekhakota": "rekhasamala",
    "shadabshaikh": "shaikhshadab",
    "srinivasgoud": "srinivaspolagoni",
    "vijaykumarbade": "vijaybade",
}


def _normalize(value):
    return "".join(ch for ch in str(value or "").lower() if ch.isalnum())


def _find_table(workbook):
    preferred = [
        ("03_All_Employees", "Office attendance dashboard"),
        ("Coding", "Employee-wise attendance"),
    ]
    candidates = preferred + [(sheet.title, "Attendance workbook") for sheet in workbook.worksheets]
    seen = set()
    for sheet_name, source_format in candidates:
        if sheet_name in seen or sheet_name not in workbook.sheetnames:
            continue
        seen.add(sheet_name)
        sheet = workbook[sheet_name]
        for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 15), values_only=True), 1):
            headers = {_normalize(value): index for index, value in enumerate(row) if value is not None}
            if REQUIRED_HEADERS.issubset(headers):
                return sheet, row_number, headers, source_format
    abort(400, message="The workbook does not contain a supported attendance table.")


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    return None


def _parse_time(value):
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    text = str(value or "").strip()
    for pattern in ("%H:%M:%S", "%H:%M", "%I:%M %p"):
        try:
            return datetime.strptime(text, pattern).time()
        except ValueError:
            pass
    return None


def _duration_minutes(value):
    if isinstance(value, timedelta):
        return max(0, round(value.total_seconds() / 60))
    if isinstance(value, (int, float)):
        return max(0, round(float(value) * 24 * 60))
    text = str(value or "").strip().lower()
    hours = re.search(r"(\d+(?:\.\d+)?)\s*h", text)
    minutes = re.search(r"(\d+(?:\.\d+)?)\s*m", text)
    if hours or minutes:
        return max(0, round(float(hours.group(1)) * 60 if hours else 0) + round(float(minutes.group(1)) if minutes else 0))
    if ":" in text:
        parts = text.split(":")
        try:
            return max(0, int(parts[0]) * 60 + int(parts[1]))
        except (ValueError, IndexError):
            return 0
    return 0


def _integer(value):
    try:
        return max(0, int(float(value or 0)))
    except (TypeError, ValueError):
        return 0


def _coding_users_by_name():
    users = (
        User.query.join(Role)
        .join(RoleType)
        .join(Project)
        .filter(
            User.is_active.is_(True),
            Project.name == "CODING",
            RoleType.code.in_(("lead", "employee")),
        )
        .all()
    )
    return {_normalize(f"{user.first_name} {user.last_name}"): user for user in users}


def import_login_hours(file_base64, source_filename, uploaded_by_id):
    try:
        payload = base64.b64decode(file_base64, validate=True)
    except (binascii.Error, ValueError):
        abort(400, message="The uploaded workbook payload is not valid base64.")
    if not payload:
        abort(400, message="The uploaded workbook is empty.")

    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception:
        abort(400, message="The uploaded file is not a readable .xlsx workbook.")

    sheet, header_row, headers, source_format = _find_table(workbook)
    users_by_name = _coding_users_by_name()
    parsed_by_key = {}
    unmatched_names = set()
    unmatched_row_count = 0

    def value(row, header):
        index = headers.get(header)
        return row[index] if index is not None and index < len(row) else None

    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        employee_name = str(value(row, "employee") or "").strip()
        attendance_date = _parse_date(value(row, "date"))
        if not employee_name or attendance_date is None:
            continue
        department = str(value(row, "department") or "").strip() or None
        if department and _normalize(department) != "coding":
            continue

        normalized_name = _normalize(employee_name)
        normalized_name = ATTENDANCE_NAME_ALIASES.get(normalized_name, normalized_name)
        user = users_by_name.get(normalized_name)
        if user is None:
            unmatched_names.add(employee_name)
            unmatched_row_count += 1
            continue

        parsed_by_key[(user.id, attendance_date)] = {
                "user": user,
                "attendance_date": attendance_date,
                "employee_name_raw": employee_name,
                "personnel_id": str(value(row, "personnelid") or "").strip() or None,
                "department": department,
                "first_in": _parse_time(value(row, "firstin")),
                "last_out": _parse_time(value(row, "lastout")),
                "total_inside_minutes": _duration_minutes(value(row, "totalinside")),
                "total_outside_minutes": _duration_minutes(value(row, "totaloutside")),
                "total_span_minutes": _duration_minutes(value(row, "totalspan")),
                "entries": _integer(value(row, "entries")),
                "exits": _integer(value(row, "exits")),
                "status": str(value(row, "status") or "").strip() or None,
                "anomalies": _integer(value(row, "anomalies")),
            }

    parsed_rows = list(parsed_by_key.values())

    batch = LoginHoursUploadBatch(
        source_filename=source_filename,
        source_format=source_format,
        uploaded_by_id=uploaded_by_id,
        row_count=len(parsed_rows) + unmatched_row_count,
        matched_count=len(parsed_rows),
        unmatched_count=unmatched_row_count,
    )
    db.session.add(batch)
    db.session.flush()

    # The newest upload wins for a user/date while the batch table retains
    # who uploaded each source and its match/drop counts.
    for data in parsed_rows:
        user = data.pop("user")
        record = LoginHourRecord.query.filter_by(
            user_id=user.id, attendance_date=data["attendance_date"]
        ).first()
        if record is None:
            record = LoginHourRecord(user_id=user.id, attendance_date=data["attendance_date"])
            db.session.add(record)
        record.batch_id = batch.id
        for key, value_to_set in data.items():
            setattr(record, key, value_to_set)

    db.session.commit()
    batch.unmatched_names = sorted(unmatched_names)
    return batch


def _readable_user_ids(actor):
    role_type = actor.role.role_type.code
    if role_type == "manager":
        return manager_team_user_ids(actor.id)
    if role_type == "lead":
        direct_report_ids = [
            user.id
            for user in User.query.join(Role).join(RoleType).filter(
                User.reports_to_id == actor.id,
                User.is_active.is_(True),
                User.project_id == actor.project_id,
                RoleType.code == "employee",
            )
        ]
        return [actor.id, *direct_report_ids]
    if role_type == "employee":
        return [actor.id]
    return []


def _login_hour_filter_options(readable_ids):
    readable_users = (
        User.query.join(Role)
        .join(RoleType)
        .filter(User.id.in_(readable_ids), User.is_active.is_(True))
        .order_by(User.first_name, User.last_name)
        .all()
    )
    leads = [user for user in readable_users if user.role.role_type.code == "lead"]
    cohorts = (
        Cohort.query.join(CohortMembership)
        .filter(CohortMembership.user_id.in_(readable_ids))
        .distinct()
        .order_by(Cohort.sequence_no.desc())
        .all()
    )
    return {
        "users": [{"id": user.id, "label": f"{user.first_name} {user.last_name}"} for user in readable_users],
        "leads": [{"id": lead.id, "label": f"{lead.first_name} {lead.last_name}"} for lead in leads],
        "cohorts": [{"id": cohort.id, "label": cohort.label} for cohort in cohorts],
    }


def list_login_hour_records(
    actor,
    page,
    page_size,
    from_date=None,
    to_date=None,
    user_id=None,
    lead_id=None,
    cohort_id=None,
):
    readable_ids = _readable_user_ids(actor)
    filter_options = _login_hour_filter_options(readable_ids)
    if user_id is not None and user_id not in readable_ids:
        abort(403, message="You cannot view login hours for that user.")

    filtered_user_ids = set(readable_ids)
    if lead_id is not None:
        allowed_lead_ids = {option["id"] for option in filter_options["leads"]}
        if lead_id not in allowed_lead_ids:
            abort(403, message="You cannot view login hours for that lead.")
        lead_team_ids = {
            row.id
            for row in User.query.filter(
                User.id.in_(readable_ids),
                db.or_(User.id == lead_id, User.reports_to_id == lead_id),
            ).all()
        }
        filtered_user_ids &= lead_team_ids

    if cohort_id is not None:
        allowed_cohort_ids = {option["id"] for option in filter_options["cohorts"]}
        if cohort_id not in allowed_cohort_ids:
            abort(403, message="You cannot view login hours for that cohort.")
        cohort_user_ids = {
            row.user_id
            for row in CohortMembership.query.filter(
                CohortMembership.cohort_id == cohort_id,
                CohortMembership.user_id.in_(readable_ids),
            ).all()
        }
        filtered_user_ids &= cohort_user_ids

    if user_id is not None:
        filtered_user_ids &= {user_id}

    query = LoginHourRecord.query.filter(LoginHourRecord.user_id.in_(filtered_user_ids))
    if from_date is not None:
        query = query.filter(LoginHourRecord.attendance_date >= from_date)
    if to_date is not None:
        query = query.filter(LoginHourRecord.attendance_date <= to_date)
    total = query.order_by(None).count()
    average_inside_minutes = query.with_entities(db.func.avg(LoginHourRecord.total_inside_minutes)).scalar()
    items = (
        query.order_by(LoginHourRecord.attendance_date.desc(), LoginHourRecord.user_id)
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    return {
        "items": items,
        "page": page,
        "page_size": page_size,
        "total": total,
        "total_pages": (total + page_size - 1) // page_size,
        "average_inside_minutes": float(average_inside_minutes) if average_inside_minutes is not None else None,
        "filter_options": filter_options,
    }
