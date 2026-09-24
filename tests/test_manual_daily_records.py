import base64
import datetime as dt
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.encryption.passwords import hash_password
from app.extensions import db
from app.manual_daily_records.models import ManualDailyRecord
from app.manual_daily_records.services import (
    MANUAL_MTD_UPLOAD_HEADERS,
    MANUAL_UPLOAD_HEADERS,
    upsert_own_record,
)
from app.roles.models import Role, RoleType
from app.users.models import User


def _get_or_create_user(email, first_name, last_name, emp_id):
    role_type = RoleType.query.filter_by(code="employee").one()
    role = Role.query.filter_by(role_type_id=role_type.id).first()
    user = User.query.filter_by(email=email).first()
    if user is None:
        user = User(
            email=email,
            first_name=first_name,
            last_name=last_name,
            emp_id=emp_id,
            role_id=role.id,
            password_hash=hash_password("test-password"),
            first_login=False,
            is_active=True,
        )
        db.session.add(user)
        db.session.commit()
    return user


def _entry(**overrides):
    entry = {
        "record_date": dt.date(2026, 9, 10),
        "production_count": 12,
        "tech_issues_downtime_hours": 1,
        "no_inventory_idle_time_hours": 0.5,
        "leave_hours": 0,
        "meeting_engagement_hours": 1.5,
    }
    entry.update(overrides)
    return entry


def _entry_body(**overrides):
    body = {
        "date": "2026-09-10",
        "productionCount": 12,
        "techIssuesDowntimeHours": 1,
        "noInventoryIdleTimeHours": 0.5,
        "leaveHours": 0,
        "meetingEngagementHours": 1.5,
    }
    body.update(overrides)
    return body


def _bulk_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "09182026"
    sheet.append(MANUAL_UPLOAD_HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return base64.b64encode(output.getvalue()).decode("ascii")


def _manager_team_member(manager, suffix="one"):
    project = manager.project
    lead_type = RoleType.query.filter_by(code="lead").one()
    employee_type = RoleType.query.filter_by(code="employee").one()
    lead_role = Role.query.filter_by(role_type_id=lead_type.id).first()
    employee_role = Role.query.filter_by(role_type_id=employee_type.id).first()
    lead = User.query.filter_by(email=f"mdr-bulk-lead-{suffix}@example.com").first()
    if lead is None:
        lead = User(
            email=f"mdr-bulk-lead-{suffix}@example.com",
            first_name="Bulk",
            last_name=f"Lead {suffix}",
            emp_id=f"TEST-MDR-BULK-LEAD-{suffix}",
            role_id=lead_role.id,
            project_id=project.id,
            reports_to_id=manager.id,
            password_hash=hash_password("test-password"),
            first_login=False,
            is_active=True,
        )
        db.session.add(lead)
        db.session.flush()
    employee = User.query.filter_by(email=f"mdr-bulk-{suffix}@example.com").first()
    if employee is None:
        employee = User(
            email=f"mdr-bulk-{suffix}@example.com",
            first_name="Bulk",
            last_name=f"Employee {suffix}",
            emp_id=f"TEST-MDR-BULK-{suffix}",
            role_id=employee_role.id,
            project_id=project.id,
            reports_to_id=lead.id,
            password_hash=hash_password("test-password"),
            first_login=False,
            is_active=True,
        )
        db.session.add(employee)
    else:
        employee.project_id = project.id
        employee.reports_to_id = lead.id
        employee.is_active = True
    db.session.commit()
    return employee


# --------------------------------------------------------------------------
# Service-level: the upsert itself, without the HTTP/crypto layer.
# --------------------------------------------------------------------------


def test_upsert_creates_new_record_as_pending():
    user = _get_or_create_user("mdr-owner1@example.com", "Owner", "One", "TEST-MDR-1")

    record = upsert_own_record(user.id, _entry())

    assert record.user_id == user.id
    assert record.record_date == dt.date(2026, 9, 10)
    assert record.production_count == 12
    assert record.status == "pending"


def test_upsert_updates_existing_record_instead_of_duplicating():
    user = _get_or_create_user("mdr-owner2@example.com", "Owner", "Two", "TEST-MDR-2")

    first = upsert_own_record(user.id, _entry(production_count=10))
    second = upsert_own_record(user.id, _entry(production_count=25))

    assert first.id == second.id
    assert ManualDailyRecord.query.filter_by(user_id=user.id, record_date=dt.date(2026, 9, 10)).count() == 1
    assert second.production_count == 25


def test_upsert_resets_status_to_pending_on_edit():
    user = _get_or_create_user("mdr-owner3@example.com", "Owner", "Three", "TEST-MDR-3")
    manager = _get_or_create_user("mdr-manager1@example.com", "Manager", "One", "TEST-MDR-MGR-1")

    record = upsert_own_record(user.id, _entry())
    record.status = "approved"
    record.reviewed_by_id = manager.id
    record.reviewed_at = dt.datetime.now(dt.timezone.utc)
    db.session.commit()

    edited = upsert_own_record(user.id, _entry(production_count=99))

    assert edited.status == "pending"
    assert edited.reviewed_by_id is None
    assert edited.reviewed_at is None
    assert edited.production_count == 99


# --------------------------------------------------------------------------
# API-level: self-entry, role gating on review, and the group filters -
# through the real encrypted request/response cycle (see conftest.ApiClient).
# --------------------------------------------------------------------------


def test_employee_can_self_enter_and_edit_own_record(api_client, employee_user):
    status, _ = api_client.login("test-employee@example.com", "test-password")
    assert status == 200

    status, body = api_client.post("/api/manual-daily-records", _entry_body())
    assert status == 200, body
    assert body["data"]["userId"] == employee_user.id
    assert body["data"]["status"] == "pending"
    assert body["data"]["productionCount"] == 12

    status, body = api_client.post("/api/manual-daily-records", _entry_body(productionCount=30))
    assert status == 200, body
    assert body["data"]["productionCount"] == 30

    status, body = api_client.get(f"/api/manual-daily-records?userId={employee_user.id}")
    assert status == 200, body
    assert len(body["data"]) == 1
    assert body["data"][0]["productionCount"] == 30


def test_program_counts_are_added_into_production_total(api_client, employee_user):
    api_client.login("test-employee@example.com", "test-password")

    status, body = api_client.post(
        "/api/manual-daily-records",
        _entry_body(productionCount=None, pvpCount=11, foundationCount=7),
    )

    assert status == 200, body
    assert body["data"]["pvpCount"] == 11
    assert body["data"]["foundationCount"] == 7
    assert body["data"]["productionCount"] == 18


def test_production_count_cannot_be_negative(api_client, employee_user):
    api_client.login("test-employee@example.com", "test-password")

    status, body = api_client.post("/api/manual-daily-records", _entry_body(productionCount=-1))
    assert status == 422, body


def test_hour_field_cannot_exceed_ten(api_client, employee_user):
    api_client.login("test-employee@example.com", "test-password")

    status, body = api_client.post("/api/manual-daily-records", _entry_body(techIssuesDowntimeHours=40))
    assert status == 422, body


def test_non_manager_cannot_approve_or_reject(api_client, employee_user):
    api_client.login("test-employee@example.com", "test-password")
    status, body = api_client.post("/api/manual-daily-records", _entry_body())
    record_id = body["data"]["id"]

    status, body = api_client.post(f"/api/manual-daily-records/{record_id}/approve")
    assert status == 403, body


def test_manager_can_approve_pending_record(api_client, employee_user, manager_user):
    api_client.login("test-employee@example.com", "test-password")
    status, body = api_client.post("/api/manual-daily-records", _entry_body())
    record_id = body["data"]["id"]

    api_client.login("test-manager@example.com", "test-password")
    status, body = api_client.post(f"/api/manual-daily-records/{record_id}/approve")
    assert status == 200, body
    assert body["data"]["status"] == "approved"
    assert body["data"]["reviewedById"] == manager_user.id

    # Already reviewed - a second decision is rejected outright.
    status, body = api_client.post(f"/api/manual-daily-records/{record_id}/reject", {"reason": "too late"})
    assert status == 409, body


def test_manager_can_reject_pending_record_with_reason(api_client, employee_user, manager_user):
    api_client.login("test-employee@example.com", "test-password")
    status, body = api_client.post("/api/manual-daily-records", _entry_body())
    record_id = body["data"]["id"]

    api_client.login("test-manager@example.com", "test-password")
    status, body = api_client.post(
        f"/api/manual-daily-records/{record_id}/reject", {"reason": "hours look wrong"}
    )
    assert status == 200, body
    assert body["data"]["status"] == "rejected"
    assert body["data"]["rejectionReason"] == "hours look wrong"


def test_query_filters_by_user_ids_and_excludes(api_client, manager_user):
    user_a = _get_or_create_user("mdr-a@example.com", "Alpha", "User", "TEST-MDR-A")
    user_b = _get_or_create_user("mdr-b@example.com", "Beta", "User", "TEST-MDR-B")
    upsert_own_record(user_a.id, _entry(record_date=dt.date(2026, 9, 11)))
    upsert_own_record(user_b.id, _entry(record_date=dt.date(2026, 9, 11)))

    api_client.login("test-manager@example.com", "test-password")

    status, body = api_client.get(f"/api/manual-daily-records?userIds={user_a.id}&userIds={user_b.id}")
    assert status == 200, body
    assert {row["userId"] for row in body["data"]} == {user_a.id, user_b.id}

    status, body = api_client.get(
        f"/api/manual-daily-records?fromDate=2026-09-11&toDate=2026-09-11&excludeUserIds={user_b.id}"
    )
    assert status == 200, body
    assert {row["userId"] for row in body["data"]} == {user_a.id}


def test_manager_downloads_manual_bulk_template(api_client, manager_user):
    api_client.login(manager_user.email, "test-password")

    response = api_client.client.get("/api/manual-daily-records/upload-template")

    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(response.data), read_only=True, data_only=True)
    assert tuple(cell.value for cell in workbook.active[1]) == MANUAL_MTD_UPLOAD_HEADERS


def _import_row(employee, date="2026-09-18", production=18):
    return {
        "userId": employee.id,
        "date": date,
        "productionCount": production,
        "techIssuesDowntimeHours": 0.5,
        "noInventoryIdleTimeHours": 0,
        "leaveHours": 0,
        "meetingEngagementHours": 0.75,
    }


def _start_manual_import(api_client, rows, suffix="one"):
    status, body = api_client.post(
        "/api/manual-daily-records/imports",
        {
            "sourceFilename": f"manual-{suffix}.xlsx",
            "fileChecksum": suffix[0] * 64,
            "totalRows": len(rows),
        },
    )
    assert status == 201, body
    return body["data"]["id"]


def test_chunked_manual_import_is_idempotent_and_preserves_identical_approval(api_client, manager_user):
    employee = _manager_team_member(manager_user, "chunked")
    rows = [_import_row(employee, "2026-09-17", 17), _import_row(employee, "2026-09-18", 18)]
    api_client.login(manager_user.email, "test-password")

    import_id = _start_manual_import(api_client, rows, "a")
    status, body = api_client.post(
        f"/api/manual-daily-records/imports/{import_id}/chunks/0",
        {"checksum": "b" * 64, "rows": rows},
    )
    assert status == 200, body
    assert body["data"]["createdCount"] == 2
    status, body = api_client.post(f"/api/manual-daily-records/imports/{import_id}/complete")
    assert status == 200, body

    record = ManualDailyRecord.query.filter_by(
        user_id=employee.id, record_date=dt.date(2026, 9, 18)
    ).one()
    record.status = "approved"
    db.session.commit()

    repeated_id = _start_manual_import(api_client, rows, "c")
    status, body = api_client.post(
        f"/api/manual-daily-records/imports/{repeated_id}/chunks/0",
        {"checksum": "d" * 64, "rows": rows},
    )
    assert status == 200, body
    assert body["data"]["unchangedCount"] == 2
    assert body["data"]["updatedCount"] == 0
    assert ManualDailyRecord.query.filter_by(user_id=employee.id).count() == 2
    assert db.session.get(ManualDailyRecord, record.id).status == "approved"


def test_chunked_manual_import_updates_only_changed_user_day(api_client, manager_user):
    employee = _manager_team_member(manager_user, "changed")
    upsert_own_record(
        employee.id,
        _entry(record_date=dt.date(2026, 9, 18), production_count=18),
    )
    api_client.login(manager_user.email, "test-password")
    rows = [_import_row(employee, production=27)]
    import_id = _start_manual_import(api_client, rows, "e")

    status, body = api_client.post(
        f"/api/manual-daily-records/imports/{import_id}/chunks/0",
        {"checksum": "f" * 64, "rows": rows},
    )

    assert status == 200, body
    assert body["data"]["updatedCount"] == 1
    record = ManualDailyRecord.query.filter_by(
        user_id=employee.id, record_date=dt.date(2026, 9, 18)
    ).one()
    assert record.production_count == 27
    assert record.status == "pending"


def test_manager_bulk_upload_creates_and_updates_team_records(api_client, manager_user):
    employee = _manager_team_member(manager_user, "success")
    payload = _bulk_workbook(
        [[employee.email, f"{employee.first_name} {employee.last_name}", 18, 0.5, 0, 0, 0.75, 0]]
    )
    api_client.login(manager_user.email, "test-password")

    status, body = api_client.post(
        "/api/manual-daily-records/bulk-upload",
        {"recordDate": "2026-09-18", "sourceFilename": "production.xlsx", "fileBase64": payload},
    )

    assert status == 201, body
    assert body["data"]["createdCount"] == 1
    assert body["data"]["importedCount"] == 1
    record = ManualDailyRecord.query.filter_by(user_id=employee.id, record_date=dt.date(2026, 9, 18)).one()
    assert record.production_count == 18
    assert record.pvp_count == 18
    assert record.foundation_count == 0

    updated_payload = _bulk_workbook(
        [[employee.email.upper(), f"{employee.first_name} {employee.last_name}", 22, 0, 1, 0, 0.25, 0]]
    )
    status, body = api_client.post(
        "/api/manual-daily-records/bulk-upload",
        {"recordDate": "2026-09-18", "sourceFilename": "production.xlsx", "fileBase64": updated_payload},
    )
    assert status == 201, body
    assert body["data"]["updatedCount"] == 1
    assert ManualDailyRecord.query.filter_by(user_id=employee.id, record_date=dt.date(2026, 9, 18)).count() == 1
    assert ManualDailyRecord.query.filter_by(user_id=employee.id, record_date=dt.date(2026, 9, 18)).one().production_count == 22


def test_bulk_upload_reports_row_errors_and_imports_nothing(api_client, manager_user):
    employee = _manager_team_member(manager_user, "invalid")
    payload = _bulk_workbook(
        [
            [employee.email, "Wrong Person", 18, 0, 0, 0, 0.75, 0],
            ["outside@example.com", "Outside Person", 10, 0, 0, 0, 0, 0],
        ]
    )
    api_client.login(manager_user.email, "test-password")

    status, body = api_client.post(
        "/api/manual-daily-records/bulk-upload",
        {"recordDate": "2026-09-18", "sourceFilename": "invalid.xlsx", "fileBase64": payload},
    )

    assert status == 422, body
    assert len(body["data"]["rowErrors"]) == 2
    assert ManualDailyRecord.query.filter_by(record_date=dt.date(2026, 9, 18)).count() == 0
